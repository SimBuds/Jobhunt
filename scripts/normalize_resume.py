"""Normalise a resume .docx after a round-trip through Word or Google Drive.

Drive is a good editing surface and a lossy export. Every download reliably
reintroduces the same defects, and each one is invisible until something
downstream breaks:

- **`docProps` metadata.** Google stamps a creator and timestamps on export.
  `render_docx._scrub_metadata` exists because recruiter-side screeners read
  these; a hand-maintained resume should carry none at all.
- **Runs with no explicit font size.** `styles.xml` here declares no default,
  so an unsized run renders at whatever the *opening application* defaults to
  — and Word and LibreOffice disagree. The contact block is the usual victim,
  showing one size on your screen and another on the reader's.
- **`pageBreakBefore` flags.** Observed 11 of them plus a hard page break in a
  single export, which pushed PROJECTS onto a third page. The page count looks
  like a content problem and is not.
- **Fragmented paragraph spacing.** Exports scatter before/after/line values
  across a document that had three consistent classes.

None of this changes a word of text, so `parse_docx` keeps working and the
verified facts are untouched — which is exactly why it goes unnoticed.

Paragraph roles are inferred from structure (bullet, heading, entry header,
stack line, contact, body) and each role gets one size and one spacing rule,
so the hierarchy the document already expresses is preserved rather than
flattened.

Run:  uv run python scripts/normalize_resume.py [path]   (default: the resume
      `resume.locate` would pick)
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parent.parent

# One size and one spacing rule per role: (space_before, space_after, line).
SIZE_PT = {"name": 16.0, "heading": 12.0}
BODY_PT = 11.0
SPACING: dict[str, tuple[float | None, float, float | None]] = {
    "name": (None, 2, None),
    "contact": (None, 2, None),
    "heading": (10, 3, None),
    "entry": (6, 2, None),
    "stack": (0, 2, None),
    "bullet": (None, 4, 1.15),
    "body": (None, 4, 1.15),
}
MARGIN_IN = 0.5


def _run_text(run) -> str:  # noqa: ANN001 - lxml element
    return "".join(n.text or "" for n in run.iter(qn("w:t")))


def classify(paras: list) -> list[str]:  # noqa: ANN001
    """Infer each paragraph's role from structure, not from its current styling
    — the styling is what we are about to overwrite."""
    first_heading = next(
        (
            i
            for i, p in enumerate(paras)
            if i > 0 and p.text.strip().isupper() and len(p.text.strip()) < 45
        ),
        1,
    )
    kinds = []
    for i, p in enumerate(paras):
        text = p.text.strip()
        if p._p.find(".//" + qn("w:numPr")) is not None:
            kinds.append("bullet")
        elif i == 0:
            kinds.append("name")
        elif text.isupper() and len(text) < 45:
            kinds.append("heading")
        elif i < first_heading:
            kinds.append("contact")
        elif text.startswith("Stack:"):
            kinds.append("stack")
        elif " | " in text:
            kinds.append("entry")
        else:
            kinds.append("body")
    return kinds


def strip_breaks(doc) -> tuple[int, int]:  # noqa: ANN001
    pbb = brk = 0
    for p in doc.paragraphs:
        p_pr = p._p.find(qn("w:pPr"))
        if p_pr is not None:
            for el in p_pr.findall(qn("w:pageBreakBefore")):
                p_pr.remove(el)
                pbb += 1
        for run in list(p._p.iter(qn("w:r"))):
            for br in run.findall(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    run.remove(br)
                    brk += 1
    return pbb, brk


def pin_fonts(doc, kinds: list[str]) -> int:  # noqa: ANN001
    """Give every run an explicit size and family. Iterates the XML rather than
    `paragraph.runs`, which silently skips runs nested inside a `w:hyperlink` —
    that omission is how the email and URL runs stayed unsized last time."""
    fixed = 0
    paras = [p for p in doc.paragraphs if p.text.strip()]
    for p, kind in zip(paras, kinds, strict=True):
        pt = SIZE_PT.get(kind, BODY_PT)
        for run in p._p.iter(qn("w:r")):
            if not _run_text(run).strip():
                continue
            r_pr = run.find(qn("w:rPr"))
            if r_pr is None:
                r_pr = OxmlElement("w:rPr")
                run.insert(0, r_pr)
            had = r_pr.find(qn("w:sz")) is not None
            for tag in ("w:sz", "w:szCs"):
                el = r_pr.find(qn(tag))
                if el is None:
                    el = OxmlElement(tag)
                    r_pr.append(el)
                el.set(qn("w:val"), str(int(pt * 2)))
            fonts = r_pr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, fonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(attr), "Calibri")
            if not had:
                fixed += 1
    return fixed


def apply_spacing(doc, kinds: list[str]) -> None:  # noqa: ANN001
    paras = [p for p in doc.paragraphs if p.text.strip()]
    for p, kind in zip(paras, kinds, strict=True):
        before, after, line = SPACING[kind]
        fmt = p.paragraph_format
        fmt.space_before = None if before is None else Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = line
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(MARGIN_IN)
        section.left_margin = section.right_margin = Inches(MARGIN_IN)


def strip_docprops(path: Path) -> list[str]:
    """Remove docProps parts and the references to them. python-docx cannot
    delete a package part, so the archive is rewritten entry by entry."""
    with zipfile.ZipFile(path) as zin:
        names = [n for n in zin.namelist() if n.startswith("docProps/")]
        if not names:
            return []
        order = [n for n in zin.namelist() if not n.startswith("docProps/")]
        keep = {n: zin.read(n) for n in order}
        infos = {i.filename: i for i in zin.infolist()}

    content_types = keep["[Content_Types].xml"].decode("utf-8")
    for part in ("/docProps/core.xml", "/docProps/app.xml", "/docProps/custom.xml"):
        pattern = rf'<Override PartName="{re.escape(part)}"[^/]*/>'
        content_types = re.sub(pattern, "", content_types)
    keep["[Content_Types].xml"] = content_types.encode("utf-8")
    rels = keep["_rels/.rels"].decode("utf-8")
    rels = re.sub(r'<Relationship[^>]*Target="docProps/[^"]*"[^>]*/>', "", rels)
    keep["_rels/.rels"] = rels.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in order:
            info = zipfile.ZipInfo(name, date_time=infos[name].date_time)
            info.compress_type = infos[name].compress_type
            zout.writestr(info, keep[name])
    return names


def page_count(path: Path) -> str:
    """Real page count via LibreOffice. The estimator in `render_docx` is tuned
    for its own output and disagrees with the real render on hand-made files."""
    if shutil.which("soffice") is None:
        return "unknown (soffice not installed)"
    out_dir = path.parent / ".normalize-tmp"
    out_dir.mkdir(exist_ok=True)
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
            capture_output=True,
            check=False,
        )
        pdf = out_dir / (path.stem + ".pdf")
        if not pdf.is_file():
            return "unknown (conversion failed)"
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        for line in info.splitlines():
            if line.startswith("Pages"):
                return line.split()[1]
        return "unknown"
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def normalize(path: Path) -> None:
    doc = Document(str(path))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    if not paras:
        print(f"{path.name}: no content")
        return
    kinds = classify(paras)

    pbb, brk = strip_breaks(doc)
    unsized = pin_fonts(doc, kinds)
    apply_spacing(doc, kinds)
    doc.save(str(path))
    props = strip_docprops(path)

    counts: dict[str, int] = {}
    for k in kinds:
        counts[k] = counts.get(k, 0) + 1
    print(f"{path.name}")
    print(f"  paragraphs        : {counts}")
    print(f"  pageBreakBefore   : removed {pbb}")
    print(f"  hard page breaks  : removed {brk}")
    print(f"  runs given a size : {unsized}")
    print(f"  docProps          : removed {len(props)} part(s)")
    print(f"  margins           : {MARGIN_IN}in all round")
    print(f"  pages             : {page_count(path)}")


def main(argv: list[str]) -> int:
    if len(argv) > 1:
        target = Path(argv[1])
    else:
        from jobhunt.resume.locate import find_baseline_resume

        target = find_baseline_resume(REPO)
    if not target.is_file():
        print(f"not found: {target}")
        return 1
    normalize(target)

    # The text is untouched, so a clean parse afterwards is the real proof.
    from jobhunt.resume.parse_docx import parse_baseline

    facts, warnings = parse_baseline(target)
    print(f"  parser warnings   : {warnings or 'none'}")
    buckets = {
        "core": facts.skills_core,
        "cms": facts.skills_cms,
        "data/devops": facts.skills_data_devops,
        "ai": facts.skills_ai,
        "projects": facts.skills_projects,
        "familiar": facts.skills_familiar,
    }
    print("  buckets           : " + " · ".join(f"{k} {len(v)}" for k, v in buckets.items()))
    empty = [k for k, v in buckets.items() if not v]
    if empty:
        print(f"  WARNING: empty bucket(s): {', '.join(empty)} — check the skill row labels")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
